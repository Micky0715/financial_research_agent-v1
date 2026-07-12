"""DOCX / PDF report export (v3 stage F).

- DOCX: python-docx，从 Markdown 报告解析标题/章节/列表/表格，嵌入 HTML 报告
  里的 base64 图表，追加来源列表、估值假设说明和免责声明。
- PDF: 优先用本机 Edge headless 从 HTML 报告打印（Windows 自带，无需额外依赖）；
  Edge 不可用或打印失败时降级——只产出 DOCX/HTML，返回 degraded 状态，不让
  导出失败影响任何上游流程。

输出目录：outputs/final_reports/docx/ 和 outputs/final_reports/pdf/（默认不入库）。
"""
import base64
import io
import re
import subprocess
from pathlib import Path
from typing import Any, Optional

from config import config
from utils.logger import logger

_DOCX_DIR = config.OUTPUT_DIR / "final_reports" / "docx"
_PDF_DIR = config.OUTPUT_DIR / "final_reports" / "pdf"

_EDGE_CANDIDATES = [
    Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
    Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
]

_DISCLAIMER = (
    "免责声明：本报告由多智能体系统基于公开网页/PDF资料与 AkShare 公开数据自动生成，"
    "估值结果为工具计算与配置假设的演示，仅供技术研究参考，不构成投资建议。"
    "数据可能存在时效性问题、抽取误差或模型生成偏差，使用者需自行核实。"
)
_ASSUMPTION_NOTE = (
    "估值假设说明：DCF 估值中的折现率、永续增长率为配置假设（见 config.py DCF_*）；"
    "基期现金流可能使用净利润代理（未扣资本开支，系统性偏乐观）；相对估值的参照系为"
    "该股票自身近一年倍数分布，非行业可比公司。"
)

_MD_TABLE_ROW = re.compile(r"^\s*\|(.+)\|\s*$")
_BASE64_IMG = re.compile(r'data:image/png;base64,([A-Za-z0-9+/=]+)')


def _md_to_docx(doc, markdown: str) -> None:
    """Minimal markdown -> docx: headings, bullet lists, tables, paragraphs."""
    lines = (markdown or "").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif _MD_TABLE_ROW.match(line):
            # collect the whole table block
            block = []
            while i < len(lines) and _MD_TABLE_ROW.match(lines[i]):
                block.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            rows = [r for r in block if not all(set(c) <= set("-: ") for c in r)]
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row[: len(rows[0])]):
                        table.cell(ri, ci).text = re.sub(r"\*\*", "", cell)
            continue
        elif line.lstrip().startswith(("- ", "* ")):
            doc.add_paragraph(re.sub(r"\*\*", "", line.lstrip()[2:]), style="List Bullet")
        elif re.match(r"^\s*\d+\.\s", line):
            doc.add_paragraph(re.sub(r"\*\*", "", re.sub(r"^\s*\d+\.\s", "", line)), style="List Number")
        else:
            doc.add_paragraph(re.sub(r"\*\*", "", line))
        i += 1


def export_docx(
    markdown: str,
    title: str,
    out_name: str,
    sources: Optional[list[dict]] = None,
    html_for_charts: str = "",
) -> dict[str, Any]:
    """Markdown 报告 -> DOCX。Returns {success, path|None, degraded, error}。"""
    try:
        from docx import Document
        from docx.shared import Inches
    except Exception as exc:  # noqa: BLE001 - python-docx missing -> degraded
        return {"success": False, "path": None, "degraded": True, "error": f"python-docx unavailable: {exc}"}

    try:
        _DOCX_DIR.mkdir(parents=True, exist_ok=True)
        doc = Document()
        doc.add_heading(title, level=0)
        _md_to_docx(doc, markdown)

        # 图表：从 HTML 报告里抠 base64 PNG 嵌入
        for b64 in _BASE64_IMG.findall(html_for_charts or "")[:4]:
            try:
                doc.add_picture(io.BytesIO(base64.b64decode(b64)), width=Inches(6.0))
            except Exception:  # noqa: BLE001 - a bad image never blocks the export
                continue

        if sources:
            doc.add_heading("来源列表", level=2)
            for s in sources:
                tier = s.get("authority_tier", "")
                doc.add_paragraph(
                    f"[{s.get('source_id')}] {s.get('title', '')[:80]} - {s.get('url', '')}"
                    + (f"（{tier}）" if tier else ""),
                    style="List Bullet",
                )

        doc.add_heading("估值假设说明", level=2)
        doc.add_paragraph(_ASSUMPTION_NOTE)
        doc.add_heading("免责声明", level=2)
        doc.add_paragraph(_DISCLAIMER)

        path = _DOCX_DIR / f"{out_name}.docx"
        doc.save(str(path))
        return {"success": True, "path": str(path), "degraded": False, "error": None}
    except Exception as exc:  # noqa: BLE001
        logger.warning(f"docx export failed: {exc}")
        return {"success": False, "path": None, "degraded": True, "error": str(exc)[:200]}


def _find_edge() -> Optional[Path]:
    for p in _EDGE_CANDIDATES:
        if p.exists():
            return p
    return None


def export_pdf(html_path: str, out_name: str) -> dict[str, Any]:
    """HTML 报告 -> PDF（Edge headless）。环境缺 Edge 或打印失败 -> degraded，不抛异常。"""
    edge = _find_edge()
    if edge is None:
        return {"success": False, "path": None, "degraded": True,
                "error": "msedge not found - PDF export degraded (DOCX/HTML only)"}
    src = Path(html_path)
    if not src.exists():
        return {"success": False, "path": None, "degraded": True, "error": f"html not found: {html_path}"}
    try:
        _PDF_DIR.mkdir(parents=True, exist_ok=True)
        out = _PDF_DIR / f"{out_name}.pdf"
        cmd = [
            str(edge), "--headless", "--disable-gpu", "--no-sandbox",
            f"--print-to-pdf={out}", "--print-to-pdf-no-header", src.resolve().as_uri(),
        ]
        proc = subprocess.run(cmd, capture_output=True, timeout=90)
        if out.exists() and out.stat().st_size > 1000:
            return {"success": True, "path": str(out), "degraded": False, "error": None}
        return {"success": False, "path": None, "degraded": True,
                "error": f"edge print produced no output (rc={proc.returncode})"}
    except Exception as exc:  # noqa: BLE001
        logger.warning(f"pdf export failed (degraded): {exc}")
        return {"success": False, "path": None, "degraded": True, "error": str(exc)[:200]}
